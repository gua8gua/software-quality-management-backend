import json
from io import BytesIO

from docx import Document
from sqlalchemy import func, select

from app.modules.tlr.models import TlrArtifact, TlrDataset, TlrFile
from app.modules.tlr.uploads import extract_text
from tests.test_tlr import SCOPE, create_run
from tests.test_tlr import api as api


async def new_project(client):
    response = await client.post(
        "/api/v1/tlr/projects",
        json={**SCOPE, "name": "质量管理测试项目", "description": "上传与追踪"},
    )
    assert response.status_code == 200, response.text


def upload_data(*, base=None, replace=False, kind="requirement"):
    data = {
        "tenant_id": SCOPE["tenant_id"],
        "version": "v1",
        "metadata": json.dumps([{"external_id": "R1", "kind": kind, "revision": "v1"}]),
        "replace_existing": str(replace).lower(),
    }
    if base:
        data["base_dataset_id"] = base
    return data


UPLOAD = f"/api/v1/tlr/projects/{SCOPE['project_id']}/upload"


async def test_project_browse_upload_detail_download_are_persistent(api):
    client, factory, _ = api
    await new_project(client)
    projects = await client.get("/api/v1/tlr/projects", params={"tenant_id": SCOPE["tenant_id"]})
    assert projects.json()["data"]["items"][0]["name"] == "质量管理测试项目"
    original = "系统应支持登录。\n".encode()
    response = await client.post(
        UPLOAD, data=upload_data(), files={"files": ("需求.md", original, "text/markdown")}
    )
    assert response.status_code == 200, response.text
    dataset = response.json()["data"]
    inventory = await client.get(f"/api/v1/tlr/datasets/{dataset['id']}/inventory", params=SCOPE)
    row = inventory.json()["data"]["items"][0]
    assert row["kind"] == "requirement" and row["original_file_id"]
    assert "content" not in row
    artifact = await client.get(f"/api/v1/tlr/artifacts/{row['id']}", params=SCOPE)
    assert artifact.json()["data"]["content"] == original.decode()
    downloaded = await client.get(f"/api/v1/tlr/artifacts/{row['id']}/download", params=SCOPE)
    assert downloaded.content == original
    assert "attachment" in downloaded.headers["content-disposition"]
    async with factory() as db:
        assert (await db.get(TlrDataset, dataset["id"])).project_id == SCOPE["project_id"]
        assert (await db.get(TlrFile, row["original_file_id"])).payload == original
    hidden = await client.get(
        f"/api/v1/tlr/artifacts/{row['id']}/download", params={**SCOPE, "tenant_id": "other"}
    )
    assert hidden.status_code == 404


async def test_upload_requires_external_project_and_file_kind(api):
    client, _, _ = api
    await new_project(client)
    response = await client.post(
        UPLOAD, data=upload_data(kind=""), files={"files": ("r.txt", b"hello")}
    )
    assert response.status_code == 422
    response = await client.post(
        UPLOAD.replace(SCOPE["project_id"], "unknown"),
        data=upload_data(),
        files={"files": ("r.txt", b"hello")},
    )
    assert response.status_code == 404
    response = await client.post(
        UPLOAD, data=upload_data(kind="automatically-infer"), files={"files": ("r.txt", b"hello")}
    )
    assert response.status_code == 422


async def test_replacement_creates_new_snapshot_and_preserves_old_file(api):
    client, _, _ = api
    await new_project(client)
    first = await client.post(
        UPLOAD, data=upload_data(), files={"files": ("r.txt", b"version one")}
    )
    first_id = first.json()["data"]["id"]
    rejected = await client.post(
        UPLOAD, data=upload_data(base=first_id), files={"files": ("r.txt", b"version two")}
    )
    assert rejected.status_code == 422
    second = await client.post(
        UPLOAD,
        data=upload_data(base=first_id, replace=True),
        files={"files": ("r.txt", b"version two")},
    )
    assert second.status_code == 200, second.text
    second_id = second.json()["data"]["id"]
    assert first_id != second_id
    for dataset_id, expected in [(first_id, "version one"), (second_id, "version two")]:
        response = await client.get(f"/api/v1/tlr/datasets/{dataset_id}/artifacts", params=SCOPE)
        assert response.json()["data"]["items"][0]["content"] == expected


async def test_invalid_batch_rolls_back_without_partial_artifacts(api):
    client, factory, _ = api
    await new_project(client)
    metadata = [
        {"external_id": "A", "kind": "code", "revision": "v1"},
        {"external_id": "B", "kind": "test_case", "revision": "v1"},
    ]
    data = {**upload_data(), "metadata": json.dumps(metadata)}
    response = await client.post(
        UPLOAD,
        data=data,
        files=[("files", ("a.py", b"print(1)")), ("files", ("bad.pdf", b"not a pdf"))],
    )
    assert response.status_code == 422
    async with factory() as db:
        for model in (TlrDataset, TlrArtifact, TlrFile):
            assert await db.scalar(select(func.count()).select_from(model)) == 0


async def test_visualization_and_evidence_scoped_to_run(api):
    client, _, _ = api
    run, _ = await create_run(client)
    assert (await client.post(f"/api/v1/tlr/runs/{run}/execute", params=SCOPE)).status_code == 200
    graph = await client.get(f"/api/v1/tlr/runs/{run}/visualization", params=SCOPE)
    data = graph.json()["data"]
    assert len(data["links"]) == 1 and len(data["candidates"]) == 4
    assert all("embedding" not in e and "content" not in e for e in data["elements"])
    candidate = data["links"][0]["evidence_candidate_ids"][0]
    detail = await client.get(f"/api/v1/tlr/runs/{run}/candidates/{candidate}", params=SCOPE)
    assert detail.json()["data"]["candidate"]["evidence"]["related"] is True
    hidden = await client.get(
        f"/api/v1/tlr/runs/{run}/visualization", params={**SCOPE, "project_id": "other"}
    )
    assert hidden.status_code == 404


def test_docx_preserves_test_case_table_text():
    doc = Document()
    doc.add_paragraph("登录测试")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "步骤：输入密码"
    table.cell(0, 1).text = "预期：登录成功"
    stream = BytesIO()
    doc.save(stream)
    content = extract_text("case.docx", stream.getvalue())
    assert content.index("登录测试") < content.index("步骤：输入密码")
    assert "预期：登录成功" in content
